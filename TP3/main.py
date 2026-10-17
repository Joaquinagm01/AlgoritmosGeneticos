"""TP3 - Problema del viajante (TSP) sobre las capitales de provincia de la República Argentina.

El programa resuelve el enunciado del tercer trabajo práctico:

  Punto 1  Justificación teórica de por qué el método exhaustivo no puede resolver el problema.
  Punto 2a Heurística del vecino más cercano partiendo de la capital que elija el usuario, con mapa.
  Punto 2b Mejor recorrido de la misma heurística probando todas las capitales como inicio.
  Punto 2c Algoritmo genético (N = 50 cromosomas, M = 200 ciclos, permutaciones de 23 ciudades,
           crossover cíclico) y comparación contra la heurística.

Las distancias salen de la tabla de la cátedra (Matriz/TablaCapitales.xlsx, en kilómetros y en línea
recta). La tabla trae 24 ciudades: las 23 capitales de provincia más la Ciudad Autónoma de Buenos Aires.
Como el enunciado pide capitales de provincia y cromosomas de 23 genes, se excluye la Ciudad Autónoma y
La Plata representa a la provincia de Buenos Aires.

Uso:
  python main.py              menú interactivo (opciones a, b y c del enunciado, más extras)
  python main.py --todo       corre todo sin menú y guarda mapas, tablas y resultados en outputs/
  python main.py --verificar  controles rápidos de la matriz y del crossover cíclico
  python main.py --help       todas las opciones
"""

from __future__ import annotations

import argparse
import csv
import itertools
import json
import math
import os
import random
import statistics
import subprocess
import sys
import time
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

# =========================
# Parámetros del enunciado y del grupo
# =========================
SEMILLA = 42             # semilla fija para que los resultados del informe se puedan reproducir
N_CROMOSOMAS = 50        # N: número de cromosomas de la población (enunciado)
N_CICLOS = 200           # M: cantidad de ciclos (enunciado)
P_CROSSOVER = 0.75       # frecuencia de crossover: a criterio del grupo (elegida con calibrar_ag)
P_MUTACION = 0.35        # frecuencia de mutación por hijo: a criterio del grupo (elegida con calibrar_ag)
OPERADOR_MUTACION = "inversion"   # "swap" (intercambio de dos genes) o "inversion" (invierte un tramo)
ELITE = 2                # cromosomas que pasan intactos a la generación siguiente
CORRIDAS_COMPARACION = 30  # corridas independientes del AG para comparar con la heurística

BASE_DIR = Path(__file__).resolve().parent
MATRIZ_XLSX = BASE_DIR / "Matriz" / "TablaCapitales.xlsx"
MAPA_JSON = BASE_DIR / "datos" / "argentina_provincias.json"
OUTPUTS_DIR = BASE_DIR / "outputs"
FIGURES_DIR = OUTPUTS_DIR / "figures"

ABRIR_MAPAS = True       # en el menú interactivo se abre cada mapa al generarlo
SEGUNDOS_POR_ANIO = 3600 * 24 * 365   # año de 365 días, como en las diapositivas de la cátedra


# =========================
# Datos: ciudades y matriz de distancias
# =========================
@dataclass(frozen=True)
class Ciudad:
    provincia: str
    capital: str
    latitud: float        # sólo se usan para dibujar el mapa
    longitud: float
    nombre_tabla: str     # cómo figura en la tabla de la cátedra


# Orden alfabético por provincia. Los genes de los cromosomas son los índices 0..22 de esta lista
# (en el informe se muestran como 1..23).
CIUDADES: List[Ciudad] = [
    Ciudad("Buenos Aires", "La Plata", -34.9214, -57.9544, "La Plata"),
    Ciudad("Catamarca", "San Fernando del Valle de Catamarca", -28.4696, -65.7795, "S.F.d.V.d. Catamarca"),
    Ciudad("Chaco", "Resistencia", -27.4516, -58.9867, "Resistencia"),
    Ciudad("Chubut", "Rawson", -43.3002, -65.1023, "Rawson"),
    Ciudad("Córdoba", "Córdoba", -31.4201, -64.1888, "Córdoba"),
    Ciudad("Corrientes", "Corrientes", -27.4692, -58.8306, "Corrientes"),
    Ciudad("Entre Ríos", "Paraná", -31.7319, -60.5238, "Paraná"),
    Ciudad("Formosa", "Formosa", -26.1850, -58.1731, "Formosa"),
    Ciudad("Jujuy", "San Salvador de Jujuy", -24.1858, -65.2995, "S.S. de Jujuy"),
    Ciudad("La Pampa", "Santa Rosa", -36.6203, -64.2906, "Santa Rosa"),
    Ciudad("La Rioja", "La Rioja", -29.4131, -66.8568, "La Rioja"),
    Ciudad("Mendoza", "Mendoza", -32.8895, -68.8458, "Mendoza"),
    Ciudad("Misiones", "Posadas", -27.3621, -55.9006, "Posadas"),
    Ciudad("Neuquén", "Neuquén", -38.9516, -68.0591, "Neuquén"),
    Ciudad("Río Negro", "Viedma", -40.8135, -63.0000, "Viedma"),
    Ciudad("Salta", "Salta", -24.7821, -65.4232, "Salta"),
    Ciudad("San Juan", "San Juan", -31.5375, -68.5364, "San Juan"),
    Ciudad("San Luis", "San Luis", -33.3017, -66.3378, "San Luis"),
    Ciudad("Santa Cruz", "Río Gallegos", -51.6230, -69.2168, "Río Gallegos"),
    Ciudad("Santa Fe", "Santa Fe", -31.6333, -60.7000, "Santa Fe"),
    Ciudad("Santiago del Estero", "Santiago del Estero", -27.7844, -64.2667, "Sgo. Del Estero"),
    Ciudad("Tierra del Fuego", "Ushuaia", -54.8019, -68.3030, "Ushuaia"),
    Ciudad("Tucumán", "San Miguel de Tucumán", -26.8083, -65.2176, "S.M. de Tucumán"),
]
N_CIUDADES = len(CIUDADES)
CABA_TABLA = "Cdad. de Bs. As."   # está en la tabla pero no es capital de provincia: se excluye

# Erratas de la tabla de la cátedra (clave normalizada -> clave correcta)
ERRATAS_TABLA = {"riogalleqos": "riogallegos"}

# Nombres con los que alguien podría querer ingresar a la Ciudad Autónoma
ALIAS_CABA = {"caba", "capitalfederal", "ciudadautonomadebuenosaires", "ciudaddebuenosaires", "cdaddebsas"}


def _clave(texto: str) -> str:
    """Texto sin tildes, en minúscula y sin espacios ni signos: sirve para comparar nombres."""
    sin_tildes = unicodedata.normalize("NFD", texto)
    sin_tildes = "".join(c for c in sin_tildes if unicodedata.category(c) != "Mn")
    return "".join(c for c in sin_tildes.lower() if c.isalnum())


CLAVES_CIUDAD = [
    {_clave(c.provincia), _clave(c.capital), _clave(c.nombre_tabla)} for c in CIUDADES
]


def leer_tabla_catedra(ruta: Path = MATRIZ_XLSX) -> Tuple[List[str], np.ndarray]:
    """Lee la tabla de la cátedra completa (24 x 24). Devuelve los nombres de fila y la matriz en km."""
    hoja = openpyxl.load_workbook(ruta, data_only=True).worksheets[0]
    nombres: List[str] = []
    fila = 2
    while hoja.cell(row=fila, column=1).value not in (None, ""):
        nombres.append(str(hoja.cell(row=fila, column=1).value).strip())
        fila += 1
    n = len(nombres)
    matriz = np.zeros((n, n), dtype=float)
    for i in range(n):
        for j in range(n):
            valor = hoja.cell(row=i + 2, column=j + 2).value
            if valor not in (None, ""):
                matriz[i, j] = float(str(valor).replace(",", "."))
    return nombres, matriz


def _indice_tabla(nombres: Sequence[str], nombre_tabla: str) -> int:
    objetivo = _clave(nombre_tabla)
    for idx, nombre in enumerate(nombres):
        clave = ERRATAS_TABLA.get(_clave(nombre), _clave(nombre))
        if clave == objetivo:
            return idx
    raise ValueError(f"La ciudad «{nombre_tabla}» no figura en la tabla de la cátedra.")


def construir_matriz() -> np.ndarray:
    """Matriz de 23 x 23 con las distancias de la tabla de la cátedra, en el orden de CIUDADES."""
    nombres, tabla = leer_tabla_catedra()
    idx = [_indice_tabla(nombres, c.nombre_tabla) for c in CIUDADES]
    return tabla[np.ix_(idx, idx)].copy()


MATRIZ = construir_matriz()


def distancia_haversine(a: Ciudad, b: Ciudad) -> float:
    """Distancia en línea recta sobre la esfera (km), sólo para contrastar la tabla de la cátedra."""
    lat1, lon1, lat2, lon2 = map(math.radians, (a.latitud, a.longitud, b.latitud, b.longitud))
    h = math.sin((lat2 - lat1) / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin((lon2 - lon1) / 2) ** 2
    return 2 * 6371.0 * math.asin(math.sqrt(h))


def validar_matriz() -> Dict[str, float]:
    """Controles de la tabla: simetría, diagonal nula y diferencia contra Haversine."""
    simetrica = bool(np.allclose(MATRIZ, MATRIZ.T))
    diagonal_nula = bool(np.allclose(np.diag(MATRIZ), 0))
    difs = []
    peor = (0.0, 0, 0)
    for i in range(N_CIUDADES):
        for j in range(i + 1, N_CIUDADES):
            h = distancia_haversine(CIUDADES[i], CIUDADES[j])
            rel = abs(MATRIZ[i, j] - h) / h
            difs.append(rel)
            if rel > peor[0]:
                peor = (rel, i, j)
    return {
        "simetrica": simetrica,
        "diagonal_nula": diagonal_nula,
        "dif_relativa_media_haversine": float(statistics.fmean(difs)),
        "dif_relativa_maxima_haversine": float(peor[0]),
        "par_mayor_diferencia": f"{CIUDADES[peor[1]].capital} - {CIUDADES[peor[2]].capital}",
    }


# =========================
# Rutas
# =========================
def ruta_distancia(ruta: Sequence[int]) -> float:
    """Longitud del circuito (incluye el regreso a la ciudad de partida)."""
    return float(sum(MATRIZ[ruta[i], ruta[(i + 1) % len(ruta)]] for i in range(len(ruta))))


def rotar(ruta: Sequence[int], inicio: int) -> List[int]:
    """Mismo circuito empezando en la ciudad `inicio`."""
    pos = list(ruta).index(inicio)
    return list(ruta[pos:]) + list(ruta[:pos])


def nombre(idx: int) -> str:
    return CIUDADES[idx].capital


def ruta_a_texto(ruta: Sequence[int]) -> str:
    nombres = [f"{nombre(i)} ({CIUDADES[i].provincia})" for i in ruta]
    nombres.append(nombres[0])
    return " -> ".join(nombres)


def mostrar_recorrido(ruta: Sequence[int]) -> None:
    """Imprime el recorrido tramo a tramo y la longitud total."""
    print(f"{'#':>3}  {'Desde':<38} {'Hasta':<38} {'km':>7}")
    total = 0.0
    for k in range(len(ruta)):
        a, b = ruta[k], ruta[(k + 1) % len(ruta)]
        d = MATRIZ[a, b]
        total += d
        print(f"{k + 1:>3}  {nombre(a):<38} {nombre(b):<38} {d:>7.0f}")
    print(f"Longitud total del recorrido: {total:.0f} km")


def resolver_ciudad(texto: str) -> int:
    """Convierte lo que escribió el usuario (número, provincia o capital) en el índice de la ciudad."""
    texto = texto.strip()
    if not texto:
        raise ValueError("No ingresaste nada.")
    if texto.isdigit():
        n = int(texto)
        if 1 <= n <= N_CIUDADES:
            return n - 1
        raise ValueError(f"El número tiene que estar entre 1 y {N_CIUDADES}.")
    clave = _clave(texto)
    if clave in ALIAS_CABA:
        raise ValueError(
            "La Ciudad Autónoma de Buenos Aires no es capital de provincia: para la provincia de Buenos Aires "
            "ingresá La Plata."
        )
    exactos = [i for i, claves in enumerate(CLAVES_CIUDAD) if clave in claves]
    if len(exactos) == 1:
        return exactos[0]
    parciales = [i for i, claves in enumerate(CLAVES_CIUDAD) if any(c.startswith(clave) for c in claves)]
    if len(parciales) == 1:
        return parciales[0]
    if len(parciales) > 1:
        opciones = ", ".join(f"{CIUDADES[i].provincia} ({nombre(i)})" for i in parciales)
        raise ValueError(f"«{texto}» es ambiguo: podría ser {opciones}.")
    raise ValueError(f"No se reconoció la provincia ni la capital «{texto}».")


# =========================
# Punto 2a y 2b: heurística del vecino más cercano
# =========================
def vecino_mas_cercano(inicio: int) -> List[int]:
    """Desde cada ciudad se va a la ciudad más cercana que todavía no se visitó (empates: la de menor índice)."""
    ruta = [inicio]
    pendientes = set(range(N_CIUDADES))
    pendientes.remove(inicio)
    actual = inicio
    while pendientes:
        siguiente = min(pendientes, key=lambda j: (MATRIZ[actual, j], j))
        ruta.append(siguiente)
        pendientes.remove(siguiente)
        actual = siguiente
    return ruta


def heuristica_todos_los_inicios() -> List[Tuple[int, List[int], float]]:
    """Resultado de la heurística desde cada capital: (inicio, ruta, longitud)."""
    resultados = []
    for inicio in range(N_CIUDADES):
        ruta = vecino_mas_cercano(inicio)
        resultados.append((inicio, ruta, ruta_distancia(ruta)))
    return resultados


def mejor_heuristica() -> Tuple[int, List[int], float]:
    """Punto 2b: el recorrido mínimo de la heurística entre todos los inicios posibles."""
    return min(heuristica_todos_los_inicios(), key=lambda r: (r[2], r[0]))


# =========================
# Punto 1: método exhaustivo (justificación teórica y medición)
# =========================
def rutas_posibles(n: int) -> Dict[str, int]:
    """Cantidad de recorridos que habría que evaluar para n ciudades."""
    return {
        "ciclos_n_factorial": math.factorial(n),               # convención de la cátedra: N! ciclos
        "inicio_fijo": math.factorial(n - 1),                  # fijando la ciudad de partida: (N-1)!
        "inicio_fijo_sin_espejos": math.factorial(n - 1) // 2,  # y sin contar el sentido inverso: (N-1)!/2
    }


def medir_fuerza_bruta(n_max: int = 10) -> List[Dict[str, float]]:
    """Resuelve por fuerza bruta (todas las permutaciones) las primeras n ciudades y mide el tiempo."""
    mediciones = []
    for n in range(5, n_max + 1):
        sub = MATRIZ[:n, :n]
        inicio_t = time.perf_counter()
        mejor = float("inf")
        cantidad = 0
        for perm in itertools.permutations(range(1, n)):
            total = sub[0, perm[0]] + sub[perm[-1], 0]
            for a, b in zip(perm, perm[1:]):
                total += sub[a, b]
            cantidad += 1
            if total < mejor:
                mejor = total
        seg = time.perf_counter() - inicio_t
        mediciones.append({"n": n, "rutas": cantidad, "segundos": seg, "rutas_por_segundo": cantidad / seg,
                           "mejor_km": float(mejor)})
    return mediciones


def anios_estimados(rutas: float, rutas_por_segundo: float) -> float:
    return rutas / rutas_por_segundo / SEGUNDOS_POR_ANIO


def analisis_exhaustivo(n_max_medicion: int = 10, verbose: bool = True) -> Dict[str, object]:
    """Punto 1: por qué el método exhaustivo no resuelve el problema con 23 capitales."""
    n = N_CIUDADES
    rutas = rutas_posibles(n)
    mediciones = medir_fuerza_bruta(n_max_medicion)
    rapidez_medida = mediciones[-1]["rutas_por_segundo"]
    datos = {
        "n": n,
        "rutas": rutas,
        "anios_catedra_1us": anios_estimados(rutas["ciclos_n_factorial"], 1e6),
        "anios_sin_espejos_1e6": anios_estimados(rutas["inicio_fijo_sin_espejos"], 1e6),
        "anios_sin_espejos_1e9": anios_estimados(rutas["inicio_fijo_sin_espejos"], 1e9),
        "anios_sin_espejos_medida": anios_estimados(rutas["inicio_fijo_sin_espejos"], rapidez_medida),
        "rutas_por_segundo_medida": rapidez_medida,
        "mediciones": mediciones,
        "held_karp_operaciones": float(n * n * 2 ** n),
    }
    if verbose:
        print("\nPunto 1: método exhaustivo para las 23 capitales")
        print(f"  Ciclos posibles (N!, criterio de la cátedra):        {rutas['ciclos_n_factorial']:.3e}")
        print(f"  Recorridos con la ciudad de partida fija ((N-1)!):    {rutas['inicio_fijo']:.3e}")
        print(f"  Sin contar el sentido inverso ((N-1)!/2):             {rutas['inicio_fijo_sin_espejos']:.3e}")
        print(f"  Tiempo a 1 µs por ciclo (supuesto de la cátedra, N!): {datos['anios_catedra_1us']:.3e} años")
        print(f"  Tiempo a 10^6 recorridos/s con (N-1)!/2:              {datos['anios_sin_espejos_1e6']:.3e} años")
        print(f"  Tiempo a 10^9 recorridos/s con (N-1)!/2:              {datos['anios_sin_espejos_1e9']:.3e} años")
        print("  Medición propia (todas las permutaciones, ciudad de partida fija):")
        for m in mediciones:
            print(f"    n = {m['n']:>2}: {m['rutas']:>9,d} recorridos en {m['segundos']:.3f} s "
                  f"({m['rutas_por_segundo']:.2e} recorridos/s)".replace(",", "."))
        print(f"  Con la rapidez medida ({rapidez_medida:.2e} recorridos/s) las 23 capitales llevarían "
              f"{datos['anios_sin_espejos_medida']:.3e} años.")
        print("  Conclusión: NO se puede resolver por búsqueda exhaustiva en un tiempo razonable.")
    return datos


def held_karp(matriz: np.ndarray, verbose: bool = False) -> Tuple[float, List[int]]:
    """Óptimo exacto por programación dinámica (Held-Karp, O(n^2 2^n)).

    No es una búsqueda exhaustiva sobre todas las permutaciones: sirve como referencia para saber
    cuánto se alejan la heurística y el algoritmo genético del mejor recorrido posible.
    Usa unos 450 MB de memoria y tarda alrededor de un minuto con 23 ciudades.
    """
    n = len(matriz)
    m = n - 1                         # ciudades 1..n-1 son los bits 0..m-1; la ciudad 0 es el inicio
    tam = 1 << m
    d = matriz.astype(np.float32)
    dm = d[1:, 1:]                    # distancias entre las m ciudades
    inf = np.float32(1e9)
    dp = np.full((tam, m), inf, dtype=np.float32)
    padre = np.full((tam, m), -1, dtype=np.int8)
    for j in range(m):
        dp[1 << j, j] = d[0, j + 1]
    masks = np.arange(tam, dtype=np.int64)
    cuenta = np.zeros(tam, dtype=np.int8)
    for b in range(m):
        cuenta += ((masks >> b) & 1).astype(np.int8)
    for s in range(2, m + 1):
        capa = masks[cuenta == s]
        for j in range(m):
            con_j = capa[((capa >> j) & 1) == 1]
            previa = con_j ^ (1 << j)
            candidatos = dp[previa] + dm[:, j]       # (cantidad, m): dp[previa, i] + d[i, j]
            mejor_i = candidatos.argmin(axis=1)
            dp[con_j, j] = candidatos[np.arange(len(con_j)), mejor_i]
            padre[con_j, j] = mejor_i.astype(np.int8)
        if verbose:
            print(f"  Held-Karp: subconjuntos de {s} ciudades listos")
    completo = tam - 1
    cierre = dp[completo] + d[1:, 0]
    ultimo = int(cierre.argmin())
    mejor = float(cierre[ultimo])
    camino = []
    mask, j = completo, ultimo
    while j != -1:
        camino.append(j + 1)
        nuevo_j = int(padre[mask, j])
        mask ^= 1 << j
        j = nuevo_j if mask else -1
    camino.append(0)
    camino.reverse()
    return mejor, camino


def optimo_exacto(forzar: bool = False, verbose: bool = True) -> Dict[str, object]:
    """Óptimo exacto de las 23 capitales; se guarda en outputs/optimo_exacto.json para no recalcularlo."""
    archivo = OUTPUTS_DIR / "optimo_exacto.json"
    if archivo.exists() and not forzar:
        return json.loads(archivo.read_text(encoding="utf-8"))
    if verbose:
        print("\nCalculando el óptimo exacto (programación dinámica de Held-Karp); tarda cerca de un minuto...")
    inicio = time.perf_counter()
    km, ruta = held_karp(MATRIZ, verbose=False)
    datos = {"distancia_km": km, "ruta": ruta, "segundos": time.perf_counter() - inicio}
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    archivo.write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")
    return datos


# =========================
# Punto 2c: algoritmo genético
# =========================
def generar_poblacion(tam: int, rng: random.Random) -> List[List[int]]:
    base = list(range(N_CIUDADES))
    poblacion = []
    for _ in range(tam):
        cromosoma = base[:]
        rng.shuffle(cromosoma)
        poblacion.append(cromosoma)
    return poblacion


def distancias_poblacion(poblacion: Sequence[Sequence[int]]) -> np.ndarray:
    pob = np.asarray(poblacion)
    return MATRIZ[pob, np.roll(pob, -1, axis=1)].sum(axis=1)


MODOS_APTITUD = ("escalada", "inversa", "complemento")


def aptitudes(distancias: np.ndarray, modo: str = "escalada") -> np.ndarray:
    """Aptitud de cada cromosoma para la ruleta. Como se minimiza la distancia, hay que invertir el orden:

    - escalada:    cuánto mejor es cada cromosoma que el peor de la generación (más 1 para que el peor conserve
                   alguna chance). Amplía el rango de la función objetivo, que es la corrección que propone el
                   apunte para la convergencia lenta; es la que usa el programa.
    - inversa:     1 / distancia.
    - complemento: suma de las distancias de la población menos la distancia del cromosoma.
    """
    if modo == "inversa":
        return 1.0 / distancias
    if modo == "complemento":
        return distancias.sum() - distancias
    return distancias.max() - distancias + 1.0


def seleccion_ruleta(poblacion: Sequence[Sequence[int]], aptitud: np.ndarray, rng: random.Random) -> List[int]:
    return list(rng.choices(poblacion, weights=aptitud, k=1)[0])


def crossover_ciclico(padre1: Sequence[int], padre2: Sequence[int]) -> Tuple[List[int], List[int]]:
    """Crossover cíclico tal como lo describe el apunte de la cátedra.

    Se toma el primer gen del padre 1 y se sigue el ciclo: se mira qué hay debajo en el padre 2, se busca ese
    valor en el padre 1 y se repite hasta volver a un valor ya usado. Esas posiciones se heredan del padre 1 y
    las restantes se completan con el padre 2 (hijo 1). El hijo 2 es el simétrico. Los dos hijos son siempre
    permutaciones válidas, por eso sirve para el viajante.
    """
    n = len(padre1)
    pos_en_padre1 = {gen: i for i, gen in enumerate(padre1)}
    en_ciclo = [False] * n
    i = 0
    while not en_ciclo[i]:
        en_ciclo[i] = True
        i = pos_en_padre1[padre2[i]]
    hijo1 = [padre1[k] if en_ciclo[k] else padre2[k] for k in range(n)]
    hijo2 = [padre2[k] if en_ciclo[k] else padre1[k] for k in range(n)]
    return hijo1, hijo2


def mutar(cromosoma: List[int], p_mutacion: float, rng: random.Random, operador: str) -> List[int]:
    if rng.random() >= p_mutacion:
        return cromosoma
    hijo = list(cromosoma)
    i, j = sorted(rng.sample(range(len(hijo)), 2))
    if operador == "swap":
        hijo[i], hijo[j] = hijo[j], hijo[i]
    else:  # inversion: invierte el orden del tramo entre dos puntos
        hijo[i:j + 1] = reversed(hijo[i:j + 1])
    return hijo


def algoritmo_genetico(
    semilla: int = SEMILLA,
    n_cromosomas: int = N_CROMOSOMAS,
    n_ciclos: int = N_CICLOS,
    p_crossover: float = P_CROSSOVER,
    p_mutacion: float = P_MUTACION,
    operador: str = OPERADOR_MUTACION,
    elite: int = ELITE,
    aptitud: str = "escalada",
) -> Dict[str, object]:
    """Una corrida del AG: ruleta + crossover cíclico + mutación + elitismo. Devuelve el historial por ciclo."""
    rng = random.Random(semilla)
    poblacion = generar_poblacion(n_cromosomas, rng)
    historial = []
    mejor_ruta: List[int] = []
    mejor_dist = float("inf")
    inicio = time.perf_counter()
    for ciclo in range(n_ciclos + 1):
        dist = distancias_poblacion(poblacion)
        idx_mejor = int(dist.argmin())
        if dist[idx_mejor] < mejor_dist:
            mejor_dist = float(dist[idx_mejor])
            mejor_ruta = list(poblacion[idx_mejor])
        historial.append({"ciclo": ciclo, "minima": float(dist.min()), "promedio": float(dist.mean()),
                          "maxima": float(dist.max())})
        if ciclo == n_ciclos:
            break
        apt = aptitudes(dist, aptitud)
        orden = np.argsort(dist)
        nueva = [list(poblacion[int(k)]) for k in orden[:elite]]
        while len(nueva) < n_cromosomas:
            padre1 = seleccion_ruleta(poblacion, apt, rng)
            padre2 = seleccion_ruleta(poblacion, apt, rng)
            if rng.random() < p_crossover:
                hijo1, hijo2 = crossover_ciclico(padre1, padre2)
            else:
                hijo1, hijo2 = padre1, padre2
            nueva.append(mutar(hijo1, p_mutacion, rng, operador))
            nueva.append(mutar(hijo2, p_mutacion, rng, operador))
        poblacion = nueva[:n_cromosomas]
    return {"mejor_ruta": mejor_ruta, "mejor_distancia": mejor_dist, "historial": historial,
            "segundos": time.perf_counter() - inicio, "semilla": semilla}


def corridas_ag(cantidad: int = CORRIDAS_COMPARACION, **kwargs) -> List[Dict[str, object]]:
    return [algoritmo_genetico(semilla=SEMILLA + k, **kwargs) for k in range(cantidad)]


def calibrar_ag(corridas: int = 30, verbose: bool = True) -> List[Dict[str, object]]:
    """Cómo se eligieron las frecuencias de crossover y de mutación (el enunciado las deja a criterio del grupo).
    Cada combinación se corre con `corridas` semillas distintas de las que se usan después para comparar."""
    filas = []
    for operador in ("swap", "inversion"):
        for pc in (0.6, 0.75, 0.9, 1.0):
            for pm in (0.05, 0.10, 0.20, 0.35, 0.50):
                dists = [algoritmo_genetico(semilla=1000 + k, p_crossover=pc, p_mutacion=pm, operador=operador)
                         ["mejor_distancia"] for k in range(corridas)]
                fila = {"operador": operador, "p_crossover": pc, "p_mutacion": pm,
                        "media_km": statistics.fmean(dists), "mejor_km": min(dists), "peor_km": max(dists),
                        "desvio_km": statistics.pstdev(dists)}
                filas.append(fila)
                if verbose:
                    print(f"  {operador:<9} PC={pc:.2f} PM={pm:.2f}: media {fila['media_km']:.0f} km, "
                          f"mejor {fila['mejor_km']:.0f} km")
    return filas


def _estadistica(distancias: Sequence[float]) -> Dict[str, float]:
    return {"media_km": statistics.fmean(distancias), "mejor_km": min(distancias), "peor_km": max(distancias),
            "desvio_km": statistics.pstdev(distancias)}


def comparar_aptitudes(corridas: int = CORRIDAS_COMPARACION) -> List[Dict[str, object]]:
    """El mismo AG con tres formas de calcular la aptitud para la ruleta (mismas semillas en las tres)."""
    filas = []
    for modo in MODOS_APTITUD:
        dists = [algoritmo_genetico(semilla=SEMILLA + k, aptitud=modo)["mejor_distancia"] for k in range(corridas)]
        filas.append({"aptitud": modo, **_estadistica(dists)})
    return filas


def estudio_ciclos(valores: Sequence[int] = (50, 100, 200, 500, 1000),
                   corridas: int = CORRIDAS_COMPARACION) -> List[Dict[str, object]]:
    """Cómo mejora el AG al aumentar la cantidad de ciclos M (el enunciado recomienda M = 200)."""
    filas = []
    for m in valores:
        dists = [algoritmo_genetico(semilla=SEMILLA + k, n_ciclos=m)["mejor_distancia"] for k in range(corridas)]
        filas.append({"ciclos": m, **_estadistica(dists)})
    return filas


# =========================
# Mapas y gráficos
# =========================
def cargar_provincias() -> List[dict]:
    if not MAPA_JSON.exists():
        return []
    return json.loads(MAPA_JSON.read_text(encoding="utf-8"))["provincias"]


NOMBRE_CORTO = {"San Fernando del Valle de Catamarca": "Catamarca", "San Miguel de Tucumán": "Tucumán",
                "San Salvador de Jujuy": "Jujuy", "Santiago del Estero": "Sgo. del Estero"}

# posiciones que se prueban para cada etiqueta: (desplazamiento x, y en puntos, alineación horizontal, vertical)
CANDIDATOS_ETIQUETA = [
    (6, 3, "left", "bottom"), (6, -3, "left", "top"), (-6, 3, "right", "bottom"), (-6, -3, "right", "top"),
    (0, 8, "center", "bottom"), (0, -8, "center", "top"), (9, 0, "left", "center"), (-9, 0, "right", "center"),
    (16, 11, "left", "bottom"), (16, -11, "left", "top"), (-16, 11, "right", "bottom"), (-16, -11, "right", "top"),
    (0, 17, "center", "bottom"), (0, -17, "center", "top"), (24, 0, "left", "center"), (-24, 0, "right", "center"),
]


def colocar_etiquetas(fig, ax, items: Sequence[Tuple[float, float, str]]) -> None:
    """Ubica cada etiqueta en la primera posición que no pisa a otra etiqueta ni a un punto del mapa."""
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    radio = 5.0 * fig.dpi / 72.0   # un punto del mapa ocupa unos 5 pt de radio
    puntos = [ax.transData.transform((x, y)) for x, y, _ in items]
    ocupado = [matplotlib.transforms.Bbox([[px - radio, py - radio], [px + radio, py + radio]]) for px, py in puntos]
    for (x, y, texto) in items:
        elegido, elegido_caja, elegido_pisado = None, None, None
        for dx, dy, ha, va in CANDIDATOS_ETIQUETA:
            t = ax.annotate(texto, (x, y), xytext=(dx, dy), textcoords="offset points", fontsize=7.6,
                            color="#1d1d1d", zorder=6, ha=ha, va=va)
            caja = t.get_window_extent(renderer).expanded(1.04, 1.12)
            pisado = sum(1 for o in ocupado if caja.overlaps(o))
            if elegido is None or pisado < elegido_pisado:
                if elegido is not None:
                    elegido.remove()
                elegido, elegido_caja, elegido_pisado = t, caja, pisado
            else:
                t.remove()
            if elegido_pisado == 0:
                break
        ocupado.append(elegido_caja)


def dibujar_base(ax, provincias: List[dict]) -> None:
    for prov in provincias:
        for pol in prov["poligonos"]:
            xs, ys = zip(*pol)
            ax.fill(xs, ys, facecolor="#eef1f5", edgecolor="#9aa5b4", linewidth=0.6, zorder=1)
    ax.set_xlim(-74.5, -53.0)
    ax.set_ylim(-56.0, -21.0)
    ax.set_aspect(1 / math.cos(math.radians(38)))
    ax.set_xlabel("Longitud")
    ax.set_ylabel("Latitud")
    ax.grid(True, linestyle=":", linewidth=0.4, alpha=0.6, zorder=0)


def guardar_mapa(ruta: Sequence[int], archivo: Path, titulo: str) -> None:
    """Mapa de la República con el recorrido, el orden de visita y la ciudad de partida."""
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(6.4, 9.0), dpi=150)
    dibujar_base(ax, cargar_provincias())
    pts = [(CIUDADES[i].longitud, CIUDADES[i].latitud) for i in ruta]
    pts.append(pts[0])
    for k in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[k], pts[k + 1]
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), zorder=3,
                    arrowprops=dict(arrowstyle="-|>", color="#1f4e79", lw=1.6, shrinkA=3, shrinkB=3,
                                    mutation_scale=9))
    items = []
    for orden, i in enumerate(ruta, start=1):
        c = CIUDADES[i]
        es_inicio = orden == 1
        ax.scatter([c.longitud], [c.latitud], s=190 if es_inicio else 34,
                   color="#2e8b57" if es_inicio else "#d1495b", edgecolor="white", linewidth=0.9, zorder=5,
                   marker="*" if es_inicio else "o")
        items.append((c.longitud, c.latitud, f"{orden}. {NOMBRE_CORTO.get(c.capital, c.capital)}"))
    ax.set_title(titulo, fontsize=11, fontweight="bold")
    fig.tight_layout()
    colocar_etiquetas(fig, ax, items)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_evolucion(historial: Sequence[dict], archivo: Path, titulo: str,
                      referencias: Optional[Dict[str, float]] = None) -> None:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    ciclos = [h["ciclo"] for h in historial]
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    ax.fill_between(ciclos, [h["minima"] for h in historial], [h["maxima"] for h in historial],
                    color="#1f4e79", alpha=0.10, label="Rango de la población (mín.-máx.)")
    ax.plot(ciclos, [h["promedio"] for h in historial], color="#d1495b", linewidth=1.8, label="Promedio de la población")
    ax.plot(ciclos, [h["minima"] for h in historial], color="#1f4e79", linewidth=2.2, label="Mejor cromosoma")
    colores = ["#2e8b57", "#b06a2c", "#533483"]
    for (texto, valor), color in zip((referencias or {}).items(), colores):
        ax.axhline(valor, linestyle="--", linewidth=1.2, color=color, label=f"{texto}: {valor:.0f} km")
    ax.set_title(titulo)
    ax.set_xlabel("Ciclo")
    ax.set_ylabel("Distancia total del recorrido (km)")
    ax.grid(True, linestyle="--", alpha=0.35)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_barras_inicios(resultados: Sequence[Tuple[int, List[int], float]], archivo: Path) -> None:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    orden = sorted(resultados, key=lambda r: r[2])
    nombres = [NOMBRE_CORTO.get(nombre(r[0]), nombre(r[0])) for r in orden]
    valores = [r[2] for r in orden]
    colores = ["#2e8b57"] + ["#1f4e79"] * (len(orden) - 2) + ["#d1495b"]
    fig, ax = plt.subplots(figsize=(8.4, 5.2))
    ax.bar(range(len(orden)), valores, color=colores)
    ax.set_xticks(range(len(orden)))
    ax.set_xticklabels(nombres, rotation=60, ha="right", fontsize=8)
    ax.set_ylim(min(valores) * 0.85, max(valores) * 1.04)
    ax.set_ylabel("Longitud del recorrido (km)")
    ax.set_title("Heurística del vecino más cercano según la capital de partida")
    ax.grid(True, axis="y", linestyle="--", alpha=0.35)
    for k, v in enumerate(valores):
        ax.text(k, v + 40, f"{v:.0f}", ha="center", fontsize=6.5, rotation=90)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_comparacion(corridas: Sequence[dict], heuristica: float, optimo: Optional[float], archivo: Path) -> None:
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    dists = [c["mejor_distancia"] for c in corridas]
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    ax.hist(dists, bins=12, color="#1f4e79", alpha=0.85, edgecolor="white")
    ax.axvline(heuristica, color="#d1495b", linestyle="--", linewidth=2, label=f"Heurística, mejor inicio: {heuristica:.0f} km")
    if optimo is not None:
        ax.axvline(optimo, color="#2e8b57", linestyle="--", linewidth=2, label=f"Óptimo exacto: {optimo:.0f} km")
    ax.axvline(statistics.fmean(dists), color="#b06a2c", linestyle=":", linewidth=2,
               label=f"Promedio del AG: {statistics.fmean(dists):.0f} km")
    ax.set_title(f"Mejor recorrido del AG en {len(dists)} corridas independientes (N=50, M=200)")
    ax.set_xlabel("Distancia total del mejor recorrido (km)")
    ax.set_ylabel("Cantidad de corridas")
    ax.legend(fontsize=8)
    ax.grid(True, linestyle="--", alpha=0.35)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_ciclos(filas: Sequence[dict], referencias: Dict[str, float], archivo: Path) -> None:
    """Distancia promedio del AG según la cantidad de ciclos, con la heurística y el óptimo como referencia."""
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    ms = [f["ciclos"] for f in filas]
    fig, ax = plt.subplots(figsize=(8.4, 4.6))
    ax.fill_between(ms, [f["mejor_km"] for f in filas], [f["peor_km"] for f in filas], color="#1f4e79", alpha=0.12,
                    label="Rango (mejor y peor corrida)")
    ax.plot(ms, [f["media_km"] for f in filas], marker="o", color="#1f4e79", linewidth=2, label="Promedio del AG")
    for (texto, valor), color in zip(referencias.items(), ("#d1495b", "#2e8b57")):
        ax.axhline(valor, linestyle="--", linewidth=1.3, color=color, label=f"{texto}: {valor:.0f} km")
    ax.axvline(N_CICLOS, color="gray", linestyle=":", linewidth=1)
    ax.text(N_CICLOS * 1.03, ax.get_ylim()[1] * 0.995, "M = 200 (enunciado)", fontsize=8, color="gray", va="top")
    ax.set_xscale("log")
    ax.set_xticks(ms)
    ax.set_xticklabels([str(m) for m in ms])
    ax.set_xlabel("Cantidad de ciclos M (N = 50)")
    ax.set_ylabel("Distancia del mejor recorrido (km)")
    ax.set_title("Efecto de la cantidad de ciclos en el algoritmo genético")
    ax.grid(True, linestyle="--", alpha=0.35)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def guardar_factorial(datos: Dict[str, object], archivo: Path) -> None:
    """Crecimiento del tiempo de la búsqueda exhaustiva con la cantidad de ciudades (escala logarítmica)."""
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    ns = list(range(5, 26))
    medida = datos["rutas_por_segundo_medida"]
    exp_medida = int(math.floor(math.log10(medida)))
    mant_medida = f"{medida / 10 ** exp_medida:.1f}".replace(".", "{,}")
    fig, ax = plt.subplots(figsize=(8.4, 4.8))
    for tasa, color, etiqueta in ((1e6, "#d1495b", r"$10^{6}$ recorridos/s (cátedra)"), (1e9, "#b06a2c", r"$10^{9}$ recorridos/s"),
                                  (medida, "#1f4e79", rf"${mant_medida}\times10^{{{exp_medida}}}$ recorridos/s (medida en esta PC)")):
        ax.plot(ns, [math.factorial(n - 1) / 2 / tasa / SEGUNDOS_POR_ANIO for n in ns], color=color,
                linewidth=2, label=etiqueta)
    ax.axhline(1, color="gray", linestyle=":", linewidth=1)
    ax.text(5.2, 1.6, "1 año", fontsize=8, color="gray")
    ax.axhline(13.8e9, color="gray", linestyle=":", linewidth=1)
    ax.text(5.2, 2.2e10, "edad del universo (13.800 millones de años)", fontsize=8, color="gray")
    ax.axvline(23, color="#2e8b57", linestyle="--", linewidth=1.5)
    ax.text(23.2, 1e-8, "23 capitales", fontsize=8, color="#2e8b57", rotation=90)
    ax.set_yscale("log")
    ax.set_xticks(range(5, 26, 2))
    ax.set_xlabel("Cantidad de ciudades")
    ax.set_ylabel("Tiempo de la búsqueda exhaustiva, (N-1)!/2 recorridos (años)")
    ax.set_title("El método exhaustivo crece factorialmente")
    ax.legend(fontsize=8, loc="lower right")
    ax.grid(True, which="both", linestyle="--", alpha=0.3)
    fig.tight_layout()
    fig.savefig(archivo, dpi=150, bbox_inches="tight")
    plt.close(fig)


def abrir_imagen(archivo: Path) -> None:
    """Abre el mapa con el visor del sistema (sólo en el menú interactivo)."""
    if not ABRIR_MAPAS:
        return
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(archivo))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(archivo)])
        else:
            subprocess.Popen(["xdg-open", str(archivo)])
    except Exception as exc:  # el mapa queda guardado igual
        print(f"(No se pudo abrir el mapa automáticamente: {exc})")


def guardar_csv(archivo: Path, filas: Sequence[dict]) -> None:
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    with archivo.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(filas[0].keys()))
        w.writeheader()
        w.writerows(filas)


# =========================
# Opciones del menú
# =========================
def listar_ciudades() -> None:
    print("\nCapitales de provincia disponibles:")
    for i, c in enumerate(CIUDADES, start=1):
        print(f"  {i:>2}. {c.provincia:<20} {c.capital}")


def opcion_a(entrada: Optional[str] = None) -> Tuple[int, List[int], float]:
    """Punto 2a: heurística desde la capital que elija el usuario, con mapa."""
    print("\n=== a) Heurística del vecino más cercano desde la capital elegida ===")
    if entrada is None:
        listar_ciudades()
    while True:
        texto = entrada if entrada is not None else input("\nIngresá la provincia o su capital (o su número): ")
        try:
            inicio = resolver_ciudad(texto)
            break
        except ValueError as exc:
            print(f"  {exc}")
            if entrada is not None:
                raise
    ruta = vecino_mas_cercano(inicio)
    distancia = ruta_distancia(ruta)
    print(f"\nCiudad de partida: {nombre(inicio)} ({CIUDADES[inicio].provincia})")
    print(f"Recorrido completo: {ruta_a_texto(ruta)}\n")
    mostrar_recorrido(ruta)
    archivo = FIGURES_DIR / f"mapa_heuristica_inicio_{_clave(nombre(inicio))}.png"
    guardar_mapa(ruta, archivo, f"Vecino más cercano desde {nombre(inicio)}: {distancia:.0f} km")
    print(f"Mapa guardado en: {archivo}")
    abrir_imagen(archivo)
    return inicio, ruta, distancia


def opcion_b() -> Tuple[int, List[int], float]:
    """Punto 2b: el recorrido mínimo de la heurística entre todos los inicios."""
    print("\n=== b) Mejor recorrido de la heurística probando todas las capitales como inicio ===")
    resultados = heuristica_todos_los_inicios()
    print(f"{'Capital de partida':<38} {'Longitud (km)':>14}")
    for inicio, _, dist in sorted(resultados, key=lambda r: r[2]):
        print(f"{nombre(inicio):<38} {dist:>14.0f}")
    inicio, ruta, distancia = min(resultados, key=lambda r: (r[2], r[0]))
    print(f"\nMejor inicio: {nombre(inicio)} ({CIUDADES[inicio].provincia})")
    print(f"Recorrido: {ruta_a_texto(ruta)}\n")
    mostrar_recorrido(ruta)
    archivo = FIGURES_DIR / "mapa_heuristica_mejor.png"
    guardar_mapa(ruta, archivo, f"Mejor recorrido de la heurística (desde {nombre(inicio)}): {distancia:.0f} km")
    guardar_barras_inicios(resultados, FIGURES_DIR / "heuristica_por_inicio.png")
    print(f"Mapa guardado en: {archivo}")
    abrir_imagen(archivo)
    return inicio, ruta, distancia


def opcion_c(semilla: int = SEMILLA) -> Dict[str, object]:
    """Punto 2c: algoritmo genético y comparación con la heurística."""
    print("\n=== c) Algoritmo genético con crossover cíclico ===")
    print(f"N = {N_CROMOSOMAS} cromosomas, M = {N_CICLOS} ciclos, permutaciones de {N_CIUDADES} ciudades, "
          f"crossover cíclico (PC = {P_CROSSOVER}), mutación por {'inversión' if OPERADOR_MUTACION == 'inversion' else 'intercambio'} "
          f"(PM = {P_MUTACION}), "
          f"elitismo de {ELITE}, semilla {semilla}.")
    resultado = algoritmo_genetico(semilla=semilla)
    inicio_h, ruta_h, dist_h = mejor_heuristica()
    ruta = rotar(resultado["mejor_ruta"], inicio_h)   # el circuito es el mismo; se lo escribe desde la misma ciudad
    distancia = resultado["mejor_distancia"]
    print(f"\nTiempo de la corrida: {resultado['segundos']:.2f} s")
    print(f"Recorrido (escrito desde {nombre(inicio_h)}): {ruta_a_texto(ruta)}\n")
    mostrar_recorrido(ruta)
    print(f"\nComparación: heurística (mejor inicio) {dist_h:.0f} km | algoritmo genético {distancia:.0f} km "
          f"| diferencia {distancia - dist_h:+.0f} km")
    archivo = FIGURES_DIR / "mapa_ag.png"
    guardar_mapa(ruta, archivo, f"Algoritmo genético: {distancia:.0f} km")
    referencias = {"Heurística, mejor inicio": dist_h}
    cache_optimo = OUTPUTS_DIR / "optimo_exacto.json"
    if cache_optimo.exists():
        referencias["Óptimo exacto"] = float(json.loads(cache_optimo.read_text(encoding="utf-8"))["distancia_km"])
    guardar_evolucion(resultado["historial"], FIGURES_DIR / "ag_evolucion.png",
                      "Evolución de la distancia en el algoritmo genético", referencias)
    print(f"Mapa guardado en: {archivo}")
    abrir_imagen(archivo)
    resultado["ruta_rotada"] = ruta
    return resultado


def opcion_exhaustivo() -> None:
    datos = analisis_exhaustivo()
    guardar_factorial(datos, FIGURES_DIR / "crecimiento_exhaustivo.png")


def opcion_optimo() -> None:
    datos = optimo_exacto(forzar=False)
    ruta = rotar(datos["ruta"], 0)
    print(f"\nÓptimo exacto (Held-Karp): {datos['distancia_km']:.0f} km, calculado en {datos['segundos']:.0f} s")
    print(f"Recorrido: {ruta_a_texto(ruta)}")


def menu() -> None:
    while True:
        print("\n=== TP3 - Problema del viajante: capitales de provincia de la República Argentina ===")
        print("  a) Heurística del vecino más cercano desde la capital que elijas (con mapa)")
        print("  b) Mejor recorrido de la heurística probando todas las capitales como inicio")
        print("  c) Algoritmo genético (N = 50, M = 200, crossover cíclico)")
        print("  d) Punto 1: justificación del método exhaustivo (cálculo y medición)")
        print("  e) Referencia: óptimo exacto por programación dinámica (tarda cerca de un minuto)")
        print("  s) Salir")
        try:
            eleccion = input("Elegí una opción: ").strip().lower()
        except (EOFError, KeyboardInterrupt):
            print("\nSaliendo...")
            return
        eleccion = {"1": "a", "2": "b", "3": "c", "4": "d", "5": "e", "0": "s"}.get(eleccion, eleccion)
        try:
            if eleccion == "a":
                opcion_a()
            elif eleccion == "b":
                opcion_b()
            elif eleccion == "c":
                opcion_c()
            elif eleccion == "d":
                opcion_exhaustivo()
            elif eleccion == "e":
                opcion_optimo()
            elif eleccion == "s":
                print("Saliendo...")
                return
            else:
                print("Opción inválida: elegí a, b, c, d, e o s.")
        except (EOFError, KeyboardInterrupt):
            print("\nOperación interrumpida.")


# =========================
# Corrida completa y verificaciones
# =========================
def correr_todo(inicio_a: str, con_optimo: bool, con_calibracion: bool) -> Dict[str, object]:
    """Resuelve todo el enunciado sin menú y guarda los resultados en outputs/ (insumo del informe)."""
    OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)
    resumen: Dict[str, object] = {"validacion_matriz": validar_matriz(), "semilla": SEMILLA}

    # Punto 1
    exh = analisis_exhaustivo()
    guardar_factorial(exh, FIGURES_DIR / "crecimiento_exhaustivo.png")
    resumen["exhaustivo"] = exh

    # Punto 2a
    idx_a, ruta_a, dist_a = opcion_a(inicio_a)
    resumen["heuristica_inicio_elegido"] = {"inicio": nombre(idx_a), "provincia": CIUDADES[idx_a].provincia,
                                            "distancia_km": dist_a, "ruta": ruta_a}

    # Punto 2b
    idx_b, ruta_b, dist_b = opcion_b()
    todos = heuristica_todos_los_inicios()
    guardar_csv(OUTPUTS_DIR / "heuristica_por_inicio.csv",
                [{"capital_inicio": nombre(i), "provincia": CIUDADES[i].provincia, "distancia_km": d,
                  "recorrido": ruta_a_texto(r)} for i, r, d in sorted(todos, key=lambda t: t[2])])
    d_todas = [t[2] for t in todos]
    resumen["heuristica_mejor"] = {"inicio": nombre(idx_b), "provincia": CIUDADES[idx_b].provincia,
                                   "distancia_km": dist_b, "ruta": ruta_b}
    resumen["heuristica_estadistica"] = {"mejor_km": min(d_todas), "peor_km": max(d_todas),
                                         "promedio_km": statistics.fmean(d_todas),
                                         "peor_inicio": nombre(max(todos, key=lambda t: t[2])[0])}

    # Referencia exacta (primero, así el gráfico del AG ya la incluye)
    if con_optimo:
        optimo_exacto()

    # Punto 2c
    ag = opcion_c()
    resumen["ag_corrida_semilla"] = {"semilla": SEMILLA, "distancia_km": ag["mejor_distancia"],
                                     "ruta": ag["ruta_rotada"], "segundos": ag["segundos"]}
    corridas = corridas_ag()
    d_ag = [c["mejor_distancia"] for c in corridas]
    guardar_csv(OUTPUTS_DIR / "ag_corridas.csv",
                [{"corrida": k + 1, "semilla": c["semilla"], "distancia_km": c["mejor_distancia"],
                  "segundos": c["segundos"]} for k, c in enumerate(corridas)])
    resumen["ag_corridas"] = {"cantidad": len(corridas), "mejor_km": min(d_ag), "peor_km": max(d_ag),
                              "promedio_km": statistics.fmean(d_ag), "desvio_km": statistics.pstdev(d_ag),
                              "mediana_km": statistics.median(d_ag),
                              "mejor_ruta": rotar(min(corridas, key=lambda c: c["mejor_distancia"])["mejor_ruta"], idx_b),
                              "corridas_que_igualan_o_superan_a_la_heuristica": sum(1 for d in d_ag if d <= dist_b),
                              "segundos_promedio": statistics.fmean(c["segundos"] for c in corridas)}

    # Referencia exacta
    optimo = None
    if con_optimo:
        datos = optimo_exacto()
        optimo = float(datos["distancia_km"])
        ruta_o = rotar(datos["ruta"], idx_b)
        resumen["optimo_exacto"] = {"distancia_km": optimo, "ruta": ruta_o, "segundos": datos["segundos"]}
        guardar_mapa(ruta_o, FIGURES_DIR / "mapa_optimo.png", f"Óptimo exacto (Held-Karp): {optimo:.0f} km")
    guardar_comparacion(corridas, dist_b, optimo, FIGURES_DIR / "ag_comparacion.png")

    # Estudios complementarios del AG (mismas semillas en todos)
    aptitud = comparar_aptitudes()
    guardar_csv(OUTPUTS_DIR / "ag_aptitudes.csv", aptitud)
    resumen["ag_aptitudes"] = aptitud
    ciclos = estudio_ciclos()
    guardar_csv(OUTPUTS_DIR / "ag_ciclos.csv", ciclos)
    resumen["ag_ciclos"] = ciclos
    refs = {"Heurística, mejor inicio": dist_b}
    if optimo is not None:
        refs["Óptimo exacto"] = optimo
    guardar_ciclos(ciclos, refs, FIGURES_DIR / "ag_ciclos.png")

    if con_calibracion:
        filas = calibrar_ag()
        guardar_csv(OUTPUTS_DIR / "calibracion_ag.csv", filas)
        resumen["calibracion_ag"] = filas

    resumen["parametros_ag"] = {"N": N_CROMOSOMAS, "M": N_CICLOS, "p_crossover": P_CROSSOVER, "p_mutacion": P_MUTACION,
                                "operador_mutacion": OPERADOR_MUTACION, "elite": ELITE}
    resumen["ciudades"] = [{"gen": i + 1, "provincia": c.provincia, "capital": c.capital} for i, c in enumerate(CIUDADES)]
    (OUTPUTS_DIR / "resumen_resultados.json").write_text(json.dumps(resumen, ensure_ascii=False, indent=2),
                                                         encoding="utf-8")
    print(f"\nResultados guardados en {OUTPUTS_DIR}")
    return resumen


def verificar() -> bool:
    """Controles rápidos: tabla de la cátedra, crossover cíclico del apunte y validez de los cromosomas."""
    ok = True

    def chequear(condicion: bool, texto: str) -> None:
        nonlocal ok
        print(("  OK    " if condicion else "  FALLA ") + texto)
        ok = ok and condicion

    print("Verificaciones:")
    v = validar_matriz()
    chequear(MATRIZ.shape == (23, 23), "la matriz es de 23 x 23 (23 capitales de provincia)")
    chequear(v["simetrica"] and v["diagonal_nula"], "la matriz es simétrica y su diagonal es cero")
    chequear(v["dif_relativa_media_haversine"] < 0.03,
             f"las distancias de la tabla coinciden con Haversine (diferencia media "
             f"{100 * v['dif_relativa_media_haversine']:.2f} %)")
    h1, h2 = crossover_ciclico([9, 8, 2, 1, 7, 4, 5, 10, 6, 3], [1, 2, 3, 4, 5, 6, 7, 8, 9, 10])
    chequear(h1 == [9, 2, 3, 1, 5, 4, 7, 8, 6, 10] and h2 == [1, 8, 2, 4, 7, 6, 5, 10, 9, 3],
             "el crossover cíclico reproduce el ejemplo del apunte de la cátedra")
    h1, h2 = crossover_ciclico([3, 4, 1, 5, 6, 2], [5, 1, 4, 6, 2, 3])
    chequear(h1 == [3, 1, 4, 5, 6, 2] and h2 == [5, 4, 1, 6, 2, 3], "el crossover cíclico resuelve el ejemplo de un final")
    rng = random.Random(7)
    validos = True
    for _ in range(500):
        a, b = generar_poblacion(2, rng)
        for hijo in crossover_ciclico(a, b):
            validos &= sorted(hijo) == list(range(N_CIUDADES))
    chequear(validos, "los hijos del crossover cíclico son siempre permutaciones válidas de las 23 ciudades")
    res = algoritmo_genetico(n_ciclos=20)
    chequear(sorted(res["mejor_ruta"]) == list(range(N_CIUDADES)) and
             abs(ruta_distancia(res["mejor_ruta"]) - res["mejor_distancia"]) < 1e-6,
             "el algoritmo genético devuelve una permutación válida y su distancia coincide")
    ruta = vecino_mas_cercano(0)
    chequear(sorted(ruta) == list(range(N_CIUDADES)), "la heurística visita cada capital una sola vez")
    coinciden = True
    for n in (6, 8, 9):
        sub = MATRIZ[:n, :n]
        exacto, _ = held_karp(sub)
        fuerza_bruta = min(sub[0, p[0]] + sum(sub[a, b] for a, b in zip(p, p[1:])) + sub[p[-1], 0]
                           for p in itertools.permutations(range(1, n)))
        coinciden &= abs(exacto - fuerza_bruta) < 1e-6
    chequear(coinciden, "Held-Karp coincide con la fuerza bruta en instancias de 6, 8 y 9 ciudades")
    return ok


def main(argv: Optional[Sequence[str]] = None) -> int:
    global ABRIR_MAPAS
    parser = argparse.ArgumentParser(description="TP3 - Problema del viajante (capitales de provincia de Argentina)")
    parser.add_argument("--todo", action="store_true", help="resuelve todo el enunciado sin menú y guarda outputs/")
    parser.add_argument("--inicio", default="Santa Fe", help="capital de partida para el punto 2a en --todo")
    parser.add_argument("--optimo", action="store_true", help="incluye el óptimo exacto (Held-Karp) en --todo")
    parser.add_argument("--calibrar", action="store_true", help="incluye la calibración de PC y PM en --todo")
    parser.add_argument("--verificar", action="store_true", help="controles rápidos de la matriz y del crossover")
    parser.add_argument("--sin-abrir", action="store_true", help="no abre los mapas al generarlos")
    args = parser.parse_args(argv)

    if args.sin_abrir or args.todo or os.environ.get("TP3_SIN_ABRIR"):
        ABRIR_MAPAS = False
    if args.verificar:
        return 0 if verificar() else 1
    if args.todo:
        random.seed(SEMILLA)
        correr_todo(args.inicio, con_optimo=args.optimo, con_calibracion=args.calibrar)
        return 0
    menu()
    return 0


if __name__ == "__main__":
    sys.exit(main())
